import io
import openpyxl
import pandas as pd
from supabase import create_client

# 1. Configura tus claves de Supabase
SUPABASE_URL = "TU_SUPABASE_URL"
SUPABASE_KEY = "TU_SUPABASE_SERVICE_ROLE_KEY" # O tu ANON KEY si RLS está deshabilitado

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

# 2. Cargar el archivo Excel con openpyxl
file_path = "inventario cias.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Inventario Bienestar']

# Mapear imágenes por número de fila (0-indexed anchor)
row_images = {}
for img in ws._images:
    anchor = img.anchor
    if hasattr(anchor, '_from') and anchor._from.row > 0:
        row_images[anchor._from.row] = img

print(f"Total imágenes detectadas en Excel: {len(row_images)}")

# 3. Recorrer las filas y registrar en Supabase
insertados = 0

for r in range(6, ws.max_row + 1):
    consecutivo = ws.cell(row=r, column=1).value
    nombre = ws.cell(row=r, column=2).value
    categoria = ws.cell(row=r, column=3).value
    cant_disponible = ws.cell(row=r, column=7).value
    estado_fisico = ws.cell(row=r, column=8).value

    if nombre and str(nombre).strip() and str(nombre).strip() != 'Nombre del elemento':
        nombre_clean = str(nombre).strip()
        cat_clean = str(categoria).strip() if categoria else 'General'
        cant_clean = cant_disponible if cant_disponible is not None else 1
        estado_clean = str(estado_fisico).strip() if estado_fisico else 'Bueno'
        placa_sena = f"SENA-BIENESTAR-{consecutivo if consecutivo else r}"

        imagen_url = None

        # Subir imagen si existe para la fila r-1
        img_obj = row_images.get(r - 1)
        if img_obj:
            try:
                img_bytes = img_obj._data()
                file_name = f"elemento-excel-{consecutivo}-{r}.png"

                # Subir al Storage
                res = supabase.storage.from_("elementos").upload(
                    file_name,
                    img_bytes,
                    file_options={"content-type": "image/png", "upsert": "true"}
                )
                
                # Obtener URL pública
                public_res = supabase.storage.from_("elementos").get_public_url(file_name)
                imagen_url = public_res
            except Exception as e:
                print(f"Aviso en imagen {nombre_clean}: {e}")

        # Insertar elemento en la base de datos
        data = {
            "nombre": nombre_clean,
            "descripcion": f"Stock disponible: {cant_clean} unidades. Estado físico: {estado_clean}",
            "categoria": cat_clean,
            "placa_sena": placa_sena,
            "imagen_url": imagen_url,
            "estado": "disponible",
            "ubication": "Almacén Bienestar al Aprendiz"
        }

        supabase.table("elementos").insert(data).execute()
        insertados += 1
        print(f"[{insertados}] Registrado: {nombre_clean} ({cat_clean}) - Foto: {'SÍ' if imagen_url else 'NO'}")

print(f"\n¡Proceso finalizado! Se subieron {insertados} elementos a Supabase.")